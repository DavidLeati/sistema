# proposal_system/document_processing/document_generator.py
from io import BytesIO
from docx import Document #
from . import docx_utils # Importação relativa

def generate_dcm_document(template_path_or_file, data_replacements, bold_placeholders, comissao_existe):
    document = Document(template_path_or_file)
    
    # 1. Preenche os placeholders
    filled_document = docx_utils.replace_placeholders_in_document(document, data_replacements, bold_placeholders)
    
    # 2. ADICIONADO: Remove linhas de tabela vazias
    filled_document = docx_utils.delete_empty_table_rows(filled_document)
    
    # 3. Processa a comissão
    filled_document = docx_utils.process_comissao_performance(filled_document, comissao_existe)
    
    # Processa paragrafos opcionais
    keep_outro_instrumento = data_replacements.get("manter_outro_instrumento", False)
    filled_document = docx_utils.process_paragraph_by_marker(
        filled_document,
        "##OUTRO_INSTRUMENTO##",
        keep_paragraph=keep_outro_instrumento
    )
    
    # 4. Reordena os índices
    filled_document = docx_utils.rescan_and_renumber_document(filled_document)

    # 5. Ajusta o layout do anexo
    filled_document = docx_utils.adjust_anexo_layout(filled_document)

    doc_io = BytesIO()
    filled_document.save(doc_io)
    doc_io.seek(0)
    return doc_io

def generate_coordenacao_document(template_path_or_file, data_replacements, bold_placeholders):
    document = Document(template_path_or_file)
    
    # 1. Preenche os placeholders
    filled_document = docx_utils.replace_placeholders_in_document(document, data_replacements, bold_placeholders)

    # 2. ADICIONADO: Remove linhas de tabela vazias
    filled_document = docx_utils.delete_empty_table_rows(filled_document)

    # 3. Ajusta o layout do anexo
    filled_document = docx_utils.adjust_anexo_layout(filled_document)

    doc_io = BytesIO()
    filled_document.save(doc_io)
    doc_io.seek(0)
    return doc_io